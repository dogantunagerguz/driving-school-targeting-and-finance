#!/usr/bin/env python3
"""Load the Power BI demo workbooks, execute SQL, and render reviewable results."""

import argparse
from datetime import date, datetime
import importlib.util
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

SQL_DIR = Path(__file__).resolve().parent
ROOT = SQL_DIR.parent
SQL_FILES = [
    "01_raw_schema.sql",
    "02_staging.sql",
    "03_marts.sql",
    "04_quality_checks.sql",
]
REPORTS = [
    ("Pipeline reconciliation", "SELECT * FROM pipeline_reconciliation"),
    ("Targeting distribution", """
        SELECT * FROM candidate_targeting_summary
        ORDER BY candidates DESC, target_segment
    """),
    ("Yearly finance KPIs", """
        SELECT * FROM yearly_finance_kpis ORDER BY calendar_year
    """),
    ("Expense-category ranking", """
        SELECT * FROM expense_category_rank ORDER BY expense_rank, expense_category
    """),
    ("Debt and certification controls", "SELECT * FROM debt_kpis"),
    ("Reported funnel arithmetic", "SELECT * FROM reported_funnel_reconciliation"),
    ("Data-quality checks", "SELECT * FROM quality_check_results ORDER BY check_name"),
]


def load_setup_module():
    path = ROOT / "scripts/setup_demo.py"
    spec = importlib.util.spec_from_file_location("driving_setup_demo", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def rows_from_workbook(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        headers = next(values)
        return [dict(zip(headers, row)) for row in values]
    finally:
        workbook.close()


def iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def load_workbooks(connection, folder):
    candidates = rows_from_workbook(folder / "Yıllık Kursiyer Listesi.xlsx")
    connection.executemany(
        "INSERT INTO raw_candidates VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        [(
            row["SN."], row["ADAY NO"], row["ADI"], row["SOYADI"], row["Cinsiyet"],
            row["GRUBU"], row["ÖĞRENİM DURUMU"], iso_date(row["D.TARİHİ"]),
            row["S.SIN."], row["Yıl"], row["Ay"], row["Grup Numarası"]
        ) for row in candidates],
    )

    accounts = rows_from_workbook(folder / "Kursiyer Genel Sınav ve Borç Listesi.xlsx")
    connection.executemany(
        "INSERT INTO raw_exam_accounts VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        [(
            row["ADAY NO"], row["TOPLAM BORÇ"], row["ÖDENEN"], row["KALAN"],
            row["SINAV HARÇ TOPLAMI"], row["ÖDENEN SINAV"], row["KALAN_1"],
            iso_date(row["SERTİFİKA VERİLİŞ TARİHİ"]), row["Yıl"], row["Ay"]
        ) for row in accounts],
    )

    income = rows_from_workbook(folder / "Gelir Listesi.xlsx")
    connection.executemany(
        "INSERT INTO raw_income VALUES (?, ?, ?, ?, ?, ?, ?)",
        [(
            row["MAKBUZ NO"], iso_date(row["TAHSİLAT TARİHİ"]), row["ADAY NO"],
            row["GRUBU"], row["TUTAR"], row["Yıl"], row["Ay"]
        ) for row in income],
    )

    expenses = rows_from_workbook(folder / "Gider Listesi.xlsx")
    connection.executemany(
        "INSERT INTO raw_expenses VALUES (?, ?, ?, ?)",
        [(
            row["SN"], iso_date(row["TARİH"]), row["İŞLEM AÇIKLAMASI"], row["TUTAR"]
        ) for row in expenses],
    )

    other_income = rows_from_workbook(folder / "Diğer Gelirler.xlsx")
    connection.executemany(
        "INSERT INTO raw_other_income VALUES (?, ?, ?, ?)",
        [(
            row["SN"], iso_date(row["TARİH"]), row["İŞLEM AÇIKLAMASI"], row["TUTAR"]
        ) for row in other_income],
    )


def build_connection():
    connection = sqlite3.connect(":memory:")
    connection.row_factory = sqlite3.Row
    connection.executescript((SQL_DIR / SQL_FILES[0]).read_text(encoding="utf-8"))
    with TemporaryDirectory() as directory:
        folder = Path(directory)
        load_setup_module().build_data(folder)
        load_workbooks(connection, folder)
    for filename in SQL_FILES[1:]:
        connection.executescript((SQL_DIR / filename).read_text(encoding="utf-8"))
    return connection


def markdown_table(rows):
    if not rows:
        return "_No rows._"
    columns = rows[0].keys()
    lines = ["| " + " | ".join(columns) + " |", "|" + "---|" * len(columns)]
    for row in rows:
        values = ["" if row[column] is None else str(row[column]) for column in columns]
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def render_results(connection):
    sections = [
        "# Executed SQL results",
        "",
        "> Generated from the same fully synthetic Excel workbooks used by the public Power BI demo. These are portfolio-demo outputs, not private or operational results.",
        "",
        "**As-of date:** 2026-09-08 for completed age and targeting.",
        "",
        "**Scope note:** the 200 / 47 / 12 funnel is arithmetic validation of a reported snapshot only; its row-level source and reporting period are not public. The separately reported eight launch-month enrolments are not combined with it.",
    ]
    for title, query in REPORTS:
        sections.extend(["", f"## {title}", "", markdown_table(connection.execute(query).fetchall())])
    return "\n".join(sections) + "\n"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--write-results", action="store_true",
                        help="Refresh sql/RESULTS.md from the executed queries")
    args = parser.parse_args()
    connection = build_connection()
    try:
        output = render_results(connection)
        print(output, end="")
        failures = connection.execute(
            "SELECT check_name, issue_count FROM quality_check_results WHERE issue_count <> 0"
        ).fetchall()
        if failures:
            raise SystemExit("SQL demo failed one or more data-quality checks.")
        if args.write_results:
            (SQL_DIR / "RESULTS.md").write_text(output, encoding="utf-8")
            print("Updated sql/RESULTS.md.")
        print("PASS: all SQL data-quality checks returned zero issues.")
    finally:
        connection.close()


if __name__ == "__main__":
    main()

