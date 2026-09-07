#!/usr/bin/env python3
"""Build a wholly synthetic, offline demo and configure its Power BI sources."""

import argparse
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
MODEL = ROOT / 'pbip/driving_mock_report.SemanticModel/definition'
PROJECT = 'pbip/driving_mock_report.pbip'
MARKER = "demo-manifest.json"


def write_book(folder, filename, sheet, columns, rows, date_columns=()):
    """Write sample rows with the exact workbook/sheet names used by the model."""
    workbook = Workbook()
    workbook.properties.creator = "Synthetic portfolio demo"
    workbook.properties.description = "Invented sample data; not business results."
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(columns)
    for row in rows:
        if len(row) != len(columns):
            raise ValueError(f"Wrong column count in {filename}")
        worksheet.append(row)
    for column in date_columns:
        index = columns.index(column) + 1
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row, index).number_format = "yyyy-mm-dd"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    workbook.save(folder / filename)
    workbook.close()
    return {"file": filename, "sheet": sheet, "rows": len(rows), "columns": columns}


def configure_model(model, folder):
    expressions = model / "expressions.tmdl"
    source = expressions.read_text(encoding="utf-8")
    # Power Query text literals escape quotes by doubling them.
    value = folder.resolve().as_posix().replace('"', '""')
    replacement = ('expression DemoDataFolder = "' + value
                   + '" meta [IsParameterQuery=true, Type="Text", IsParameterQueryRequired=true]')
    source, count = re.subn(r'^expression DemoDataFolder = .*$',
                            lambda match: replacement, source, flags=re.MULTILINE)
    if count != 1:
        raise ValueError("Expected one DemoDataFolder parameter in expressions.tmdl")
    expressions.write_text(source, encoding="utf-8")


def build_data(folder):
    trainee_columns = ["SN.", "ADAY NO", "ADI", "SOYADI", "Cinsiyet", "GRUBU", "ÖĞRENİM DURUMU",
                       "D.TARİHİ", "S.SIN.", "Yıl", "Ay", "Grup Numarası", "1. TEL", "2. TEL"]
    exam_columns = ["SN", "ADAY NO", "GRUBU", "ADI SOYADI", "TOPLAM BORÇ", "ÖDENEN", "KALAN",
                    "SINAV HARÇ TOPLAMI", "ÖDENEN SINAV", "KALAN_1", "SERTİFİKA VERİLİŞ TARİHİ",
                    "Grup Numarası", "Yıl", "Ay"]
    income_columns = ["SN", "MAKBUZ NO", "TAHSİLAT TARİHİ", "ADAY NO", "GRUBU", "ADI SOYADI",
                      "TUTAR", "TC NO", "Yıl", "Ay", "Grup Numarası"]
    finance_columns = ["SN", "TARİH", "EVRAK NO", "HESAP NO", "İŞLEM AÇIKLAMASI", "TUTAR"]
    trainees, exams, income = [], [], []
    for i in range(1, 49):
        candidate = 1000 + i
        year, month = 2020 + (i - 1) % 6, 1 + (i - 1) % 12
        group = f"{100 + i}/{year}-{month:02} Group 1"
        first, last = "DEMO", f"CANDIDATE {i:03}"
        certificate = ["B", "B", "C", "A1", "A2", "D"][(i - 1) % 6]
        birth = date([2004, 1990, 1985, 2007, 2002, 1975][(i - 1) % 6], 1, 15)
        trainees.append([i, candidate, first, last, "K" if i % 7 == 0 else "E", group,
                         "Lise", birth, certificate, year, month, 1, f"DEMO-PHONE-{i:03}", ""])
        fee = 8000 + (year - 2020) * 3000
        paid = fee if i % 4 else fee // 2
        issued = date(year, month, 20) if i % 5 else None
        exams.append([i, candidate, group, f"{first} {last}", fee, paid, fee - paid,
                      1000, 1000, 0, issued, 1, year, month])
        income.append([i, 2000 + i, date(year, month, 10), candidate, group,
                       f"{first} {last}", paid, 0, year, month, 1])
    expenses, other = [], []
    descriptions = ["YAKIT", "ARAÇ BAKIM", "PERSONEL (SÜRÜCÜ KURSU)", "ÇAY",
                    "BANKA", "KİRA", "KİTAP", "DEMİRBAŞ", "DEMO DİĞER"]
    for year in range(2020, 2026):
        for month in range(1, 13):
            for category, description in enumerate(descriptions):
                i = len(expenses) + 1
                expenses.append([i, date(year, month, 15), 3000 + i, 0,
                                 description, 150 + category * 30 + (year - 2020) * 50])
            i = len(other) + 1
            other.append([i, date(year, month, 16), 5000 + i, 0, "DEMO EK GELİR", 500])
    return [
        write_book(folder, "Yıllık Kursiyer Listesi.xlsx", "Sayfa1", trainee_columns, trainees, ["D.TARİHİ"]),
        write_book(folder, "Kursiyer Genel Sınav ve Borç Listesi.xlsx", "Sheet1", exam_columns, exams,
                   ["SERTİFİKA VERİLİŞ TARİHİ"]),
        write_book(folder, "Gelir Listesi.xlsx", "Sayfa1", income_columns, income, ["TAHSİLAT TARİHİ"]),
        write_book(folder, "Gider Listesi.xlsx", "Sheet1", finance_columns, expenses, ["TARİH"]),
        write_book(folder, "Diğer Gelirler.xlsx", "Sheet1", finance_columns, other, ["TARİH"]),
    ]



def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=ROOT / "demo-data",
                        help="Output folder for generated workbooks (default: repo/demo-data)")
    args = parser.parse_args()
    folder = args.data_dir.expanduser().resolve()
    marker = folder / MARKER
    if folder.exists() and any(folder.glob("*.xlsx")):
        if not marker.exists():
            parser.error("Output contains existing workbooks; choose a new --data-dir.")
        if json.loads(marker.read_text(encoding="utf-8")).get("generator") != PROJECT:
            parser.error("Output belongs to a different demo; choose a new --data-dir.")
    folder.mkdir(parents=True, exist_ok=True)
    files = build_data(folder)
    marker.write_text(json.dumps({
        "generator": PROJECT,
        "data_kind": "fully synthetic; no private inputs or API calls",
        "reference_year": 2026,
        "files": files,
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    configure_model(MODEL, folder)
    print(f"Created {len(files)} synthetic workbooks in {folder}")
    print(f"Configured DemoDataFolder. Open {ROOT / PROJECT} in Power BI Desktop and Refresh.")


if __name__ == "__main__":
    main()
